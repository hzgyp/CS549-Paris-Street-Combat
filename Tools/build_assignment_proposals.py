"""Build the two current course proposals as DOCX and Markdown sources.

Run with the Codex managed Python runtime. Render DOCX with the documents
skill renderer, inspect every page, and copy the verified PDFs alongside them.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Docs/Proposal"
IMAGE = OUT / "Visuals/paris-mission-flowchart.png"
MOCKUP_IMAGE = OUT / "Visuals/paris-six-character-concept.png"
FAB = "https://www.fab.com/listings/dae418da-1969-444a-821c-c1f30a3f21b6"
EPIC = "https://dev.epicgames.com/documentation/en-us/unreal-engine/"
SOURCES = [
    ("France Liberation asset", FAB),
    ("Animation Notifies", EPIC + "animation-notifies-in-unreal-engine"),
    ("Collision", EPIC + "collision-in-unreal-engine---overview"),
    ("Navigation", EPIC + "navigation-system-in-unreal-engine"),
    ("Behavior Trees", EPIC + "behavior-tree-in-unreal-engine---overview"),
    ("AI Perception", EPIC + "ai-perception-in-unreal-engine"),
    ("Event Dispatchers", EPIC + "event-dispatchers-in-unreal-engine"),
]

CONCEPT = [
    "We propose a single-player first-person squad mission through a connected part of Meshingun Studio's WW2 - France Liberation city, set during the August 1944 liberation period. Our earlier Normandy mission was too complex in scope because it combined sea, beach, and fortified terrain with different movement, combat, and environmental interactions. Reusing the existing city lets us focus on character interaction and NPC design. We will survey the level in Unreal and choose linked streets with an approach, a defended objective, an exit, and an alternative route. The actual geometry will determine the mission area and traversal time.",
    "The initial configuration has six characters: one Allied player, two Allied NPCs, and three German NPCs. This is a development starting point, not a final population cap. We may add enemy groups and separately configured patrol/search routes after testing pacing, navigation, and performance. Allies accompany the player and regroup after corners. Enemies patrol, guard, investigate observed positions, and reposition during combat. Shared behavior logic supports different roles and routes without requiring a new AI system for every soldier.",
    "Intermediate objectives guide the mission: reach a rally point, clear its assigned enemy group, then reach an endpoint. Additional arrival or clearance stages can be configured later. Clearing a stage requires its registered enemy group to be defeated; walking outside an area does not count as defeat. The objective display advances when the current condition passes. Player death causes failure, and restart restores the configured roster and all objectives. Progress between stages preserves casualties and ammunition.",
    "The visual goal is coherent movement and combat across the environment. Running and aiming should blend naturally, reload events should agree with ammunition, and walls should block movement and gunfire. Allies must negotiate narrow passages without teleporting or overlapping, while enemies lose sight behind buildings and search only their last observed target positions. Fixed daylight and static cover keep these interactions readable. The mission uses selected connected routes within the larger city; it does not require every building to be enterable.",
    "Licensed soldier/rifle assets and compatible clips provide visual content, while Unreal supplies rendering and physical simulation. We will integrate these resources with student-built gameplay and evaluate the encounter through repeatable tests and a packaged Windows build. Multiplayer, driving, ocean simulation, unrestricted destruction, and detailed character modeling remain outside scope."
]

def xml(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), str(v))
    return el

def document(subtitle):
    d = Document()
    for style in d.styles:
        for border in style._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    s = d.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = Inches(.58)
    s.bottom_margin = Inches(.58)
    s.left_margin = Inches(.7)
    s.right_margin = Inches(.7)
    s.header_distance = Inches(.2)
    s.footer_distance = Inches(.25)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        st = d.styles[name]
        st.font.name = "Arial"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.05
    d.styles["Title"].font.size = Pt(23)
    d.styles["Title"].font.bold = True
    d.styles["Title"].paragraph_format.space_after = Pt(4)
    d.styles["Subtitle"].font.size = Pt(11)
    d.styles["Subtitle"].font.italic = False
    d.styles["Subtitle"].paragraph_format.space_after = Pt(10)
    for name in ["Heading 1", "Heading 2"]:
        st = d.styles[name]
        st.font.size = Pt(12 if name == "Heading 1" else 10.5)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    d.styles["Caption"].font.size = Pt(9)
    d.styles["Caption"].font.italic = False
    d.core_properties.title = "Paris Street Combat " + subtitle
    d.core_properties.author = "Yupu Guo; Yuqi Pu; Jingdi Wu"
    d.core_properties.subject = "CS549 course project proposal"
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("CS549  |  ")
    r.font.size = Pt(8)
    fld = xml("w:fldSimple", **{"w:instr": "PAGE"})
    p._p.append(fld)
    d.add_paragraph("Paris Street Combat", "Title")
    d.add_paragraph(subtitle, "Subtitle")
    return d

def para(d, text, style=None):
    p = d.add_paragraph(style=style)
    # Double asterisks delimit authored emphasis.
    for i, part in enumerate(text.split("**")):
        r = p.add_run(part)
        r.bold = bool(i % 2)
    return p

def table(d, headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = xml("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(xml("w:" + side, **{"w:val":"single", "w:sz":"4", "w:color":"D9D9D9"}))
    props.append(borders)
    margins = xml("w:tblCellMar")
    for side in ["top", "bottom"]:
        margins.append(xml("w:" + side, **{"w:w":"75", "w:type":"dxa"}))
    for side in ["left", "right"]:
        margins.append(xml("w:" + side, **{"w:w":"100", "w:type":"dxa"}))
    props.append(margins)
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    for ri, values in enumerate([headers] + rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(xml("w:cantSplit"))
        if ri == 0:
            row._tr.get_or_add_trPr().append(xml("w:tblHeader"))
        for ci, (cell, txt) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(xml("w:shd", **{"w:fill":"E5EAF0" if ri == 0 else "FFFFFF"}))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            if ci == 1 and headers[ci] == "NetID":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt)
            r.font.size = Pt(10)
            r.bold = ri == 0
    return t

def hyperlink(p, label, url):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = xml("w:hyperlink", **{"r:id": rid})
    r = xml("w:r")
    props = xml("w:rPr")
    props.append(xml("w:color", **{"w:val":"1F4D70"}))
    props.append(xml("w:sz", **{"w:val":"17"}))
    r.append(props)
    txt = xml("w:t")
    txt.text = label
    r.append(txt)
    h.append(r)
    p._p.append(h)

def sources(d, short=False, include_local=True, label="References"):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    r.font.size = Pt(8.5)
    selected = SOURCES[:1] if short else SOURCES
    for i, (label, url) in enumerate(selected):
        if i:
            r = p.add_run("  |  ")
            r.font.size = Pt(8.5)
        hyperlink(p, label if short else f"[{i + 1}] {label}", url)
    if include_local:
        p = d.add_paragraph("Local sources: Assignment 1 and Assignment 2; vendor documentation, printed pp. 9, 14, 23 and 70-75.")
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.size = Pt(8.5)

def picture(d, path, width, description):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    for prop in p._p.xpath(".//wp:docPr"):
        prop.set("descr", description)

def mockup(d, width=6.3):
    picture(d, MOCKUP_IMAGE, width, "AI-generated first-person street concept based on official France Liberation appearance references: player rifle, two Allied NPCs and three German NPCs around a junction with streets continuing into the city. The initial six-character roster is expandable; geography is conceptual.")
    para(d, "AI-generated concept (OpenAI ImageGen), informed by official Meshingun Studio France Liberation references. Initial roster; illustrative layout, not implemented gameplay. [1]", "Caption")

def flowchart(d):
    picture(d, IMAGE, 7.1, "Mission flow: start with configured roster, reach rally A, clear registered enemy group B, reach end C, succeed. Player death at any active stage causes failure; full restart restores roster and objectives. Unmet conditions keep the current stage active.")

def pillar_work_table(d, emphasize_pillars=False, include_proof=True):
    headers = ["Pillar", "Specific work", "Implementation route", "Planned proof"]
    rows = [
        ["Animation", "Run/aim transitions; fire, reload, hit and death actions.", "Retarget purchased clips; Blend Spaces + Montages. Guarded Notifies commit ammo once per reload ID; cancel stale actions. [2]", "20 interrupted action cycles; no duplicate ammo or illegal firing."],
        ["Collision Detection", "Character/wall contact; bullet hits on cover and soldiers.", "Capsules + camera-aim and muzzle-clearance/obstruction traces. First blocker controls damage; allies block shots without friendly damage. [3]", "12 fixtures at 30/60/120 FPS; compare camera-only shots at corners."],
        ["Pathfinding & Navigation", "Routes between goals; ally follow/regroup and bottleneck handling.", "Student A* at surveyed junctions; NavMesh path lengths as costs, Euclidean heuristic. MoveTo executes legs; reserve distinct destinations. [4]", "10 A*/Dijkstra pairs; equal costs, blocked paths, no persistent stalls."],
        ["NPC AI", "Guard/patrol, sight-driven combat, bounded search, return to role.", "Shared Behavior Tree; per-NPC team, role, group, patrol route and search zone. Private Blackboard; search near last seen position. [5, 6]", "10 AI/goal cases; faction, occlusion, search expiry, death and regroup."],
    ]
    widths = [1.05, 1.65, 2.85, 1.55]
    if not include_proof:
        headers = headers[:3]
        rows = [row[:3] for row in rows]
        widths = [1.35, 2.15, 3.6]
    t = table(d, headers, rows, widths)
    if emphasize_pillars:
        for row in t.rows[1:]:
            for paragraph in row.cells[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

def save(d, name):
    path = OUT / (name + ".docx")
    d.save(path)
    # Body-order Markdown source allows simple review without an Office reader.
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    lines = []
    for el in d.element.body:
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                if el.xpath(".//w:drawing"):
                    for blip in el.xpath(".//a:blip"):
                        part = d.part.related_parts[blip.get(qn("r:embed"))]
                        is_mockup = part.blob == MOCKUP_IMAGE.read_bytes()
                        label = "AI-generated street concept using official environment references" if is_mockup else "AI assisted staged mission flowchart"
                        filename = "paris-six-character-concept.png" if is_mockup else "paris-mission-flowchart.png"
                        lines.append(f"![{label}](Visuals/{filename})")
                continue
            prefix = "# " if p.style.name == "Title" else "## " if p.style.name.startswith("Heading") else ""
            lines.append(prefix + text)
        elif el.tag == qn("w:tbl"):
            t = Table(el, d)
            rows = []
            for i, row in enumerate(t.rows):
                rows.append("| " + " | ".join(c.text.replace("|", "\\|") for c in row.cells) + " |")
                if i == 0:
                    rows.append("| " + " | ".join("---" for _ in row.cells) + " |")
            lines.append("\n".join(rows))
    lines.append("\n" + "\n".join(f"- [{a}]({b})" for a, b in SOURCES))
    (OUT / (name + ".md")).write_text("\n\n".join(lines) + "\n", encoding="utf-8")
    print(path)

def assignment1():
    d = document("Assignment 1  Group Formation Pillars and Concept")
    d.add_heading("1. Group formation", 1)
    table(d, ["Group member", "NetID", "Proposed implementation responsibility"], [
        ["Yupu Guo (Group Leader)", "yg745", "Collision, gunplay and integration"],
        ["Yuqi Pu", "yp549", "Animation and character integration"],
        ["Jingdi Wu", "jw2046", "Navigation and NPC AI; environment setup"],
    ], [2.3, .7, 4.1])
    para(d, "Yupu submits for the group. Implementation roles will be agreed at kickoff, with integration support across pillars.", "Caption")
    d.add_heading("2. Project concept summary", 1)
    for text in CONCEPT:
        para(d, text)
    d.add_page_break()
    d.add_heading("3. Street concept and staged mission", 1)
    mockup(d)
    flowchart(d)
    para(d, "Initial roster: 1 player + 2 allies + 3 enemies. Add finite groups/stages after testing. Objective markers are not checkpoint saves; this flow shows logic, not map geometry.", "Caption")
    d.add_heading("4. Selected pillars and implementation plan", 1)
    pillar_work_table(d, emphasize_pillars=True, include_proof=False)
    para(d, "**Integration and physics support.** A Blueprint manager advances Reach/Clear stages from overlap/death events. Reuse compatible rigs/clips, CharacterMovement and engine physics; no custom dynamics solver.")
    sources(d, include_local=False)
    assert 250 <= len(re.findall(r"\S+", " ".join(CONCEPT))) <= 400
    save(d, "CS549_Assignment1_Proposal")

def assignment2():
    d = document("Assignment 2  Product Requirements Technical Specification and MVP")
    # Fit the added concept visual by tightening paragraph gaps, not body type.
    d.styles["Normal"].paragraph_format.space_after = Pt(3)
    for name in ["Heading 1", "Heading 2"]:
        d.styles[name].paragraph_format.space_before = Pt(6)
        d.styles[name].paragraph_format.space_after = Pt(3)
    para(d, "**Yupu Guo (yg745), Group Leader  |  Yuqi Pu (yp549)  |  Jingdi Wu (jw2046)**", "Caption")
    d.add_heading("1. Product requirements (PRD)", 1)
    para(d, "**1.1. Problem and audience.** For PC FPS players and CS549 reviewers, make animation, gunfire, cover and NPC decisions agree. Reuse France Liberation to simplify the varied terrain and interactions of the earlier Normandy mission.")
    para(d, "**1.2. Experience.** Traverse connected streets through Reach/Clear objectives and an alternate approach. Begin with 1 Allied player, 2 Allied NPCs and 3 German NPCs; six is not a final cap. Add finite groups/stages after pacing, navigation and performance checks. Casualties persist between stages.")
    d.add_heading("1.3. User stories", 2)
    for s in [
        "As a player, I want synchronized run/aim/fire/reload actions so that motion and ammunition agree.",
        "As a player, I want walls and soldiers to block shots so that cover has consistent consequences.",
        "As a player, I want allies to follow and enemies to patrol/search so that city traversal affects combat.",
        "As a player, I want intermediate objectives so that I know where to go and which group to defeat.",
    ]:
        p = para(d, s)
        p.paragraph_format.space_after = Pt(2)
    d.add_heading("1.4. Feature priorities (MoSCoW)", 2)
    table(d, ["MoSCoW", "Feature commitment"], [
        ["Must", "Initial six; configurable Reach/Clear stages and NPC routes/search zones; ally follow/regroup; one rifle; four pillars; move/aim/fire/reload/damage; health/ammo/objective UI; win/fail/reset; Windows build."],
        ["Should", "Additional enemy groups/stages after validation; crouch if supported; simple impact/footstep audio."],
        ["Could", "Checkpoint saves; hearing; exposure-weighted routes; ragdolls; improved hand IK."],
        ["Won't", "Landing/ocean, multiplayer, driving, broad interiors, destructible buildings, custom detailed soldiers, complex squad commands, infinite waves, runtime LLMs."],
    ], [.8, 6.3])
    d.add_heading("2. Technical specification", 1)
    para(d, "**2.1. Stack.** UE5; Blueprint visual scripting, C++ only if needed. Enhanced Input, UMG, Animation Blueprints, IK Retargeter, NavMesh, AIController/Behavior Trees, AI Perception and Niagara. Structs/Data Assets configure stages/NPCs; Git/LFS; no external runtime library.")
    para(d, "**2.2. Dependencies.** Original asset: UE5.6; working copy: 5.8. Pin a tested version; retain ChaosVehiclesPlugin. Purchase compatible soldier/rifle rigs and clips separately. Survey routes, collision, sightlines and NavMesh; use fixed daylight/static cover. [1]")
    para(d, "**2.3. AI strategy and performance.** ChatGPT/Codex and ImageGen assist offline planning/code/visuals; no runtime AI service. Target: 60 FPS at 1080p on i9-12900F / RTX 3080 / 32 GB Windows PC. Profile mean/p95 frame times, roster and active NPC counts; unmeasured.")
    mockup(d, width=4.6)
    d.add_page_break()
    d.add_heading("2.4. Pillars and implementation", 2)
    pillar_work_table(d, emphasize_pillars=True, include_proof=False)
    para(d, "**2.5. Shared integration.** Player/NPC weapons share ammo and hit rules; tracers are cosmetic. Configure purchased rigs, collision and clips; reuse CharacterMovement and engine physics.")
    d.add_heading("2.6. Mission flow and objective control", 2)
    flowchart(d)
    para(d, "Initial MVP: 1 player + 2 allies + 3 enemies. Add configured groups/stages after validation; objective markers do not imply checkpoint saves.", "Caption")
    para(d, "A Blueprint mission manager evaluates Reach/Clear stages from overlap/death events. Reach checks the player. Clear requires a registered, nonempty assigned group with zero survivors; leaving the area does not count. Deduplicate events, include earlier kills, and recheck state on stage activation. Full restart rejects stale callbacks and restores the configured roster. [7]")
    d.add_heading("3. Narrow vertical slice (MVP)", 1)
    para(d, "**3.1. Core slice and hardest feature.** Deliver a polished, playable three-stage mission with the initial roster, one rifle and an alternate approach. The hardest feature is keeping squad navigation, combat and objective progression consistent through the full mission.")
    para(d, "**3.2. Midterm demonstration.** Provide a packaged Windows build, gameplay recording and debug logs. Show allies following/regrouping, enemies losing sight and searching, walls blocking shots, and interrupted reloads updating ammo correctly. Compare A*/Dijkstra route costs. Demonstrate ordered objectives, persistent casualties and full restart, including early kills and duplicate events. Record frame times and path failures before expanding the roster.")
    para(d, "**3.3. MVP exclusions.** Cut additional groups/stages, checkpoint saves, advanced cover tactics, physical bullets and citywide simulation. Keep the core mission functional before expanding content.")
    sources(d, include_local=False, label="4. References")
    save(d, "CS549_Assignment2_Proposal")

if __name__ == "__main__":
    assignment1()
    assignment2()
    print("Assignment 1 concept summary words:", len(re.findall(r"\S+", " ".join(CONCEPT))))
